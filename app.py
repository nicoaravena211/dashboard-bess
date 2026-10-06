import calendar
import os
import time
import unicodedata
from datetime import datetime
from io import BytesIO

import pandas as pd
import plotly.graph_objects as go
import requests
import streamlit as st
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from streamlit_autorefresh import st_autorefresh

# =========================================================
# CONFIGURACION GENERAL
# =========================================================
st.set_page_config(page_title="Dashboard de Operación", page_icon="⚡", layout="wide", initial_sidebar_state="expanded")

COLOR_PRIMARIO = "#1565C0"
COLOR_SECUNDARIO = "#65BDEB"
COLOR_BUDGET = "#0B3D91"
COLOR_BUDGET_PPA = "#2EAD5B"
COLOR_BUDGET_PROYECTO = "#F05A28"
COLOR_TARJETA = "#FFFFFF"
COLOR_TEXTO = "#172B4D"
COLOR_TEXTO_SECUNDARIO = "#5E6C84"
COLOR_BORDE = "#DFE1E6"
COLOR_REJILLA = "rgba(94,108,132,0.18)"

CACHE_EXCEL_SEGUNDOS = 12 * 60 * 60
REFRESCO_MS = 12 * 60 * 60 * 1000
SEGUNDOS_ROTACION_GRAFICOS = 10
DIAS_CARRUSEL = 30
INCLUIR_DIA_ACTUAL_CARRUSEL = True

st_autorefresh(interval=REFRESCO_MS, key="actualizacion_dashboard_12h")

st.markdown("""
<style>
.stApp { background-color:#F5F7FA; }
.block-container { padding-top:1.4rem; padding-bottom:3rem; max-width:1600px; }
h1 { color:#172B4D; font-size:2.2rem; font-weight:760; letter-spacing:-.025em; margin-bottom:.2rem; }
h2 { color:#172B4D; font-size:1.55rem; font-weight:700; margin-top:1.8rem; padding-bottom:.45rem; border-bottom:1px solid #DFE1E6; }
h3 { color:#172B4D; font-weight:650; }
div[data-testid="stCaptionContainer"] { color:#5E6C84; }
section[data-testid="stSidebar"] { background-color:#FFFFFF; border-right:1px solid #DFE1E6; }
div[data-testid="stMetric"] { background:#FFFFFF; border:1px solid #DFE1E6; border-radius:12px; padding:15px 17px; box-shadow:0 2px 7px rgba(23,43,77,.06); min-height:112px; }
div[data-testid="stMetricLabel"] { color:#5E6C84; font-size:.88rem; font-weight:600; }
div[data-testid="stMetricValue"] { color:#172B4D; font-size:1.5rem; font-weight:750; }
div[data-testid="stPlotlyChart"] { background:#FFFFFF; border:1px solid #DFE1E6; border-radius:12px; padding:8px; box-shadow:0 2px 7px rgba(23,43,77,.05); }
div[data-testid="stExpander"] { background:#FFFFFF; border:1px solid #DFE1E6; border-radius:10px; }
div[data-baseweb="select"] > div { background:#FFFFFF; border-color:#B3BAC5; border-radius:8px; }
div[data-testid="stDataFrame"] { background:#FFFFFF; border:1px solid #DFE1E6; border-radius:10px; overflow:hidden; }
div[data-testid="stAlert"] { border-radius:10px; }
hr { border-color:#DFE1E6; margin:1.5rem 0; }
</style>
""", unsafe_allow_html=True)

try:
    API_KEY_PRMTE = st.secrets["API_KEY_PRMTE"]
    API_KEY_CMG = st.secrets["API_KEY_CMG"]
except KeyError:
    st.error("No se encontraron las credenciales API en Streamlit Secrets.")
    st.info("Configura API_KEY_PRMTE y API_KEY_CMG desde Manage app > Settings > Secrets.")
    st.stop()

# =========================================================
# CONFIGURACION OPERACIONAL
# =========================================================
ANIO_OPERACIONAL = 2026
FECHA_INICIO = f"{ANIO_OPERACIONAL}-01"
hoy = pd.Timestamp.now()
if hoy.year == ANIO_OPERACIONAL:
    FECHA_FIN = hoy.strftime("%Y-%m")
elif hoy.year > ANIO_OPERACIONAL:
    FECHA_FIN = f"{ANIO_OPERACIONAL}-12"
else:
    FECHA_FIN = f"{ANIO_OPERACIONAL}-01"

PRECIO_PPA = 83.29
TIPO_CMG = "PRELIMINAR"
BARRA_CMG_BESS = "M.ELENA_______220"
ARCHIVO_EXCEL = "budgets.xlsx"
HOJA_BUDGETS_BESS = "Budgets"
HOJA_CONSOLIDADO_BESS = "Consolidado"
HOJA_BUDGETS_PROYECTOS = "Budgets_Proyectos"
HOJA_CONSOLIDADO_PROYECTOS = "Consolidado_Proyectos"
URL_PRMTE = "https://medidas.api.coordinador.cl/medidas-v2/measurement"
URL_CMG = "https://sipub.api.coordinador.cl/costo-marginal-real/v4/findByDate"

PROYECTOS = {
    "MARIA ELENA PFV": [{"mpid":"MARELENA_220_JT1_GSS", "canal":3}],
    "SAN PEDRO III": [{"mpid":"SOLRJAMA_220_JT1_RUC", "canal":3}, {"mpid":"SOLRJAMA_220_JT2_RUC", "canal":3}],
    "DOÑA CARMEN": [{"mpid":"CDNCARMN_220_J1_ECM", "canal":3}],
    "LA QUINTA": [{"mpid":"CABRERO_023_PMGD7_QTC", "canal":3}],
    "LA PERLA": [{"mpid":"LSANGLES_013_PMGD7_LAP", "canal":3}],
    "LA HUERTA": [{"mpid":"PARRONAL_013_PMGD2_HUE", "canal":3}],
    "CHACAICO": [{"mpid":"LSANGLES_013_PMGD8_CCC", "canal":3}],
    "SANCLEMENTE": [{"mpid":"SCLMENTE_013_PMGD4_CFL", "canal":3}],
    "TRILALEO": [{"mpid":"CHOLGUAN_015_PMGD6_YTL", "canal":3}],
    "COLLANCO": [{"mpid":"CNSTUCON_023_PMGD4_OAO", "canal":3}],
}
BESS_MPIDS = ["MARELENA_023_E1_GSS", "MARELENA_023_E7_GSS", "MARELENA_023_E8_GSS", "MARELENA_023_E11_GSS"]
PROYECTOS_PMGD = ["LA QUINTA", "LA PERLA", "LA HUERTA", "CHACAICO", "SANCLEMENTE", "TRILALEO", "COLLANCO"]
ORDEN_GRAFICOS = ["MARIA ELENA PFV", "SAN PEDRO III", "DOÑA CARMEN", "CONSOLIDADO PMGDS", "LA QUINTA", "LA PERLA", "LA HUERTA", "CHACAICO", "SANCLEMENTE", "TRILALEO", "COLLANCO"]
ORDEN_CARRUSEL = ["MARIA ELENA BESS"] + list(PROYECTOS.keys())
ORDEN_DESCARGABLE_MENSUAL = ["LA QUINTA", "LA PERLA", "LA HUERTA", "CHACAICO", "SANCLEMENTE", "TRILALEO", "COLLANCO", "DOÑA CARMEN", "MARIA ELENA PFV", "SAN PEDRO III"]
MESES_CORTOS = {1:"Ene",2:"Feb",3:"Mar",4:"Abr",5:"May",6:"Jun",7:"Jul",8:"Ago",9:"Sep",10:"Oct",11:"Nov",12:"Dic"}
MESES_COMPLETOS = {1:"Enero",2:"Febrero",3:"Marzo",4:"Abril",5:"Mayo",6:"Junio",7:"Julio",8:"Agosto",9:"Septiembre",10:"Octubre",11:"Noviembre",12:"Diciembre"}

# =========================================================
# FUNCIONES VISUALES Y GENERALES
# =========================================================
def aplicar_estilo_grafico(figura, altura, leyenda_abajo=False, mostrar_leyenda=True, margen_inferior=55):
    leyenda = dict(orientation="h", yanchor="top", y=-.20, xanchor="center", x=.5, traceorder="normal") if leyenda_abajo else dict(orientation="h", yanchor="bottom", y=1.02, xanchor="left", x=0)
    if leyenda_abajo:
        margen_inferior = max(margen_inferior, 110)
    figura.update_layout(height=altura, paper_bgcolor=COLOR_TARJETA, plot_bgcolor=COLOR_TARJETA, font=dict(family="Arial, sans-serif", color=COLOR_TEXTO), title=dict(font=dict(size=17, color=COLOR_TEXTO), x=.01, xanchor="left"), legend=leyenda, showlegend=mostrar_leyenda, hovermode="x unified", margin=dict(l=65,r=30,t=70,b=margen_inferior))
    figura.update_xaxes(showgrid=False, showline=True, linecolor=COLOR_BORDE, tickfont=dict(color=COLOR_TEXTO_SECUNDARIO), title_font=dict(color=COLOR_TEXTO_SECUNDARIO))
    figura.update_yaxes(showgrid=True, gridcolor=COLOR_REJILLA, gridwidth=1, showline=False, zeroline=True, zerolinecolor=COLOR_BORDE, tickfont=dict(color=COLOR_TEXTO_SECUNDARIO), title_font=dict(color=COLOR_TEXTO_SECUNDARIO))
    return figura


def quitar_tildes(texto):
    return "".join(c for c in unicodedata.normalize("NFD", str(texto)) if unicodedata.category(c) != "Mn")


def normalizar_nombre_columna(nombre):
    nombre = quitar_tildes(nombre).strip().replace(" ", "_").replace("-", "_")
    while "__" in nombre:
        nombre = nombre.replace("__", "_")
    return nombre.lower()


def normalizar_columnas_excel(df):
    eq = {"ano":"Anio","anio":"Anio","mes":"Mes","proyecto":"Proyecto","budget_generacion_mwh":"Budget_Generacion_MWh","budget_ppa_mwh":"Budget_PPA_MWh","energia_ppa_mwh":"Energia_PPA_MWh","energia_spot_mwh":"Energia_Spot_MWh","ingreso_ppa_usd":"Ingreso_PPA_USD","ingreso_spot_usd":"Ingreso_Spot_USD","generacion_mwh":"Generacion_MWh"}
    return df.rename(columns={c:eq.get(normalizar_nombre_columna(c), str(c).strip()) for c in df.columns})


def convertir_numerico(serie):
    if pd.api.types.is_numeric_dtype(serie):
        return pd.to_numeric(serie, errors="coerce")
    return pd.to_numeric(serie.astype(str).str.strip().str.replace(" ", "", regex=False).str.replace(",", ".", regex=False), errors="coerce")


def parse_fecha_safe(serie):
    return pd.to_datetime(serie.astype(str).str[:19], errors="coerce")


def generar_periodos(fecha_inicio, fecha_fin):
    inicio, fin = datetime.strptime(fecha_inicio, "%Y-%m"), datetime.strptime(fecha_fin, "%Y-%m")
    periodos, actual = [], inicio
    while actual <= fin:
        periodos.append(actual.strftime("%Y%m"))
        actual = actual.replace(year=actual.year+1, month=1) if actual.month == 12 else actual.replace(month=actual.month+1)
    return periodos


def obtener_rango_fechas_carrusel(cantidad_dias, incluir_dia_actual=True):
    hoy_n = pd.Timestamp.now().normalize()
    fin = hoy_n if incluir_dia_actual else hoy_n - pd.Timedelta(days=1)
    return fin - pd.Timedelta(days=cantidad_dias - 1), fin


def periodo_a_texto(periodo):
    return f"{MESES_COMPLETOS[int(periodo[4:6])]} {int(periodo[:4])}"


PERIODOS_DISPONIBLES = generar_periodos(FECHA_INICIO, FECHA_FIN)
FECHA_INICIO_CARRUSEL, FECHA_FIN_CARRUSEL = obtener_rango_fechas_carrusel(DIAS_CARRUSEL, INCLUIR_DIA_ACTUAL_CARRUSEL)
PERIODOS_CARRUSEL = [p.strftime("%Y%m") for p in pd.period_range(FECHA_INICIO_CARRUSEL, FECHA_FIN_CARRUSEL, freq="M")]

@st.cache_resource
def obtener_hora_inicio_aplicacion():
    return datetime.now()
HORA_INICIO_APLICACION = obtener_hora_inicio_aplicacion()

# =========================================================
# LECTURA EXCEL
# =========================================================
def leer_hoja_excel(hoja, requeridas):
    if not os.path.exists(ARCHIVO_EXCEL):
        raise FileNotFoundError(f"No se encontró {ARCHIVO_EXCEL}.")
    df = normalizar_columnas_excel(pd.read_excel(ARCHIVO_EXCEL, sheet_name=hoja, engine="openpyxl"))
    faltantes = [c for c in requeridas if c not in df.columns]
    if faltantes:
        raise ValueError(f"Faltan columnas en {hoja}: " + ", ".join(faltantes))
    return df[requeridas].copy()


@st.cache_data(ttl=CACHE_EXCEL_SEGUNDOS, show_spinner=False)
def cargar_budgets_bess():
    req = ["Anio","Mes","Budget_Generacion_MWh","Budget_PPA_MWh"]
    df = leer_hoja_excel(HOJA_BUDGETS_BESS, req)
    for c in req: df[c] = convertir_numerico(df[c])
    df = df.dropna(subset=req); df[["Anio","Mes"]] = df[["Anio","Mes"]].astype(int)
    if df.empty or not df["Mes"].between(1,12).all() or df.duplicated(["Anio","Mes"]).any():
        raise ValueError("La hoja Budgets no contiene información válida o tiene duplicados.")
    return df


@st.cache_data(ttl=CACHE_EXCEL_SEGUNDOS, show_spinner=False)
def cargar_consolidado_bess():
    req = ["Anio","Mes","Energia_PPA_MWh","Energia_Spot_MWh","Ingreso_PPA_USD","Ingreso_Spot_USD"]
    df = leer_hoja_excel(HOJA_CONSOLIDADO_BESS, req)
    for c in req: df[c] = convertir_numerico(df[c])
    df = df.dropna(subset=req); df[["Anio","Mes"]] = df[["Anio","Mes"]].astype(int)
    if df.duplicated(["Anio","Mes"]).any(): raise ValueError("Hay meses duplicados en Consolidado.")
    df["Energia_Total_MWh"] = df["Energia_PPA_MWh"] + df["Energia_Spot_MWh"]
    df["Ingreso_Total_USD"] = df["Ingreso_PPA_USD"] + df["Ingreso_Spot_USD"]
    df["Fuente"] = "Excel consolidado"
    return df


@st.cache_data(ttl=CACHE_EXCEL_SEGUNDOS, show_spinner=False)
def cargar_budgets_proyectos():
    req = ["Anio","Mes","Proyecto","Budget_Generacion_MWh"]
    df = leer_hoja_excel(HOJA_BUDGETS_PROYECTOS, req)
    for c in ["Anio","Mes","Budget_Generacion_MWh"]: df[c] = convertir_numerico(df[c])
    df["Proyecto"] = df["Proyecto"].astype(str).str.strip().str.upper(); df = df.dropna(subset=req); df[["Anio","Mes"]] = df[["Anio","Mes"]].astype(int)
    desconocidos = sorted(set(df["Proyecto"]) - set(PROYECTOS))
    if desconocidos or df.duplicated(["Anio","Mes","Proyecto"]).any(): raise ValueError("Budgets_Proyectos contiene proyectos no reconocidos o duplicados.")
    return df


@st.cache_data(ttl=CACHE_EXCEL_SEGUNDOS, show_spinner=False)
def cargar_consolidado_proyectos():
    req = ["Anio","Mes","Proyecto","Generacion_MWh"]
    df = leer_hoja_excel(HOJA_CONSOLIDADO_PROYECTOS, req)
    for c in ["Anio","Mes","Generacion_MWh"]: df[c] = convertir_numerico(df[c])
    df["Proyecto"] = df["Proyecto"].astype(str).str.strip().str.upper(); df = df.dropna(subset=req); df[["Anio","Mes"]] = df[["Anio","Mes"]].astype(int)
    desconocidos = sorted(set(df["Proyecto"]) - set(PROYECTOS))
    if desconocidos or df.duplicated(["Anio","Mes","Proyecto"]).any() or (df["Generacion_MWh"] < 0).any(): raise ValueError("Consolidado_Proyectos contiene registros inválidos.")
    df["Fuente"] = "Excel consolidado"
    return df

# =========================================================
# API
# =========================================================
def obtener_datos_prmte(periodo, mpid, canal):
    params = {"channelId":canal,"measurePointId":mpid,"period":periodo,"user_key":API_KEY_PRMTE}
    ultimo_error = None
    for intento in range(3):
        try:
            r = requests.get(URL_PRMTE, params=params, timeout=60)
            if r.status_code == 200:
                data = r.json()
                if isinstance(data, list) and data: return data, None
                ultimo_error = "respuesta sin registros"
            elif r.status_code == 429: ultimo_error = "límite temporal de solicitudes (HTTP 429)"
            elif r.status_code >= 500: ultimo_error = f"error temporal del servidor (HTTP {r.status_code})"
            else: ultimo_error = f"respuesta HTTP {r.status_code}"
        except requests.Timeout: ultimo_error = "tiempo de espera agotado"
        except requests.ConnectionError: ultimo_error = "error de conexión"
        except requests.RequestException as error: ultimo_error = f"error de comunicación ({type(error).__name__})"
        except ValueError: ultimo_error = "respuesta JSON inválida"
        if intento < 2: time.sleep(3 * (intento + 1))
    return None, ultimo_error


def extraer_mediciones(datos, canal):
    registros, campo = [], f"channel{canal}"
    if datos is None: return registros
    for bloque in datos:
        for m in bloque.get("measurement", []): registros.append({"Fecha":m.get("dateRange"),"Energia_kWh":m.get(campo)})
    return registros


def descargar_cmg_mes(periodo):
    anio, mes = int(periodo[:4]), int(periodo[4:6]); inicio = pd.Timestamp(anio,mes,1); fin = pd.Timestamp(anio,mes,calendar.monthrange(anio,mes)[1]); hoy_n = pd.Timestamp.now().normalize()
    if anio == hoy_n.year and mes == hoy_n.month: fin = min(fin,hoy_n)
    registros, pagina = [], 0
    while True:
        params = {"startDate":inicio.strftime("%Y-%m-%d"),"endDate":fin.strftime("%Y-%m-%d"),"page":pagina,"limit":5000,"type":TIPO_CMG,"bar_transf":BARRA_CMG_BESS,"user_key":API_KEY_CMG}
        try:
            r = requests.get(URL_CMG, params=params, timeout=60)
            if r.status_code != 200: return None
            datos = r.json().get("data", [])
        except (requests.RequestException, ValueError): return None
        if not datos: break
        registros.extend(datos)
        if len(datos) < 5000: break
        pagina += 1
    if not registros: return None
    df = pd.DataFrame(registros); req = ["fecha","hra","min","cmg_usd_mwh_"]
    if not all(c in df.columns for c in req): return None
    df["hra"] = convertir_numerico(df["hra"]).fillna(0); df["min"] = convertir_numerico(df["min"]).fillna(0); df["CMG_USD_MWh"] = convertir_numerico(df["cmg_usd_mwh_"]).fillna(0)
    df["Fecha"] = pd.to_datetime(df["fecha"], errors="coerce") + pd.to_timedelta(df["hra"],unit="h") + pd.to_timedelta(df["min"],unit="m")
    return df[["Fecha","CMG_USD_MWh"]].dropna(subset=["Fecha"]).groupby("Fecha",as_index=False).agg(CMG_USD_MWh=("CMG_USD_MWh","mean"))


def obtener_pendientes_proyectos(df_consolidado):
    existentes = {(str(f.Proyecto),f"{int(f.Anio):04d}{int(f.Mes):02d}") for f in df_consolidado.itertuples()}
    return {p:[per for per in PERIODOS_DISPONIBLES if (p,per) not in existentes] for p in PROYECTOS}


@st.cache_data(show_spinner=False)
def descargar_proyectos_pendientes(pendientes):
    resultados, advertencias = [], []
    for proyecto, periodos in pendientes.items():
        for periodo in periodos:
            registros, correctos = [], 0
            for conf in PROYECTOS[proyecto]:
                datos, error = obtener_datos_prmte(periodo+"012345",conf["mpid"],conf["canal"]); regs = extraer_mediciones(datos,conf["canal"])
                if datos is None or not regs:
                    advertencias.append(f"{proyecto}, {periodo_a_texto(periodo)}, {conf['mpid']}: {error or 'sin mediciones'}"); continue
                correctos += 1; registros.extend(regs)
            if correctos != len(PROYECTOS[proyecto]) or not registros: continue
            df = pd.DataFrame(registros); df["Fecha"] = parse_fecha_safe(df["Fecha"]); df["Energia_kWh"] = convertir_numerico(df["Energia_kWh"]).fillna(0)
            df = df.dropna(subset=["Fecha"]).groupby("Fecha",as_index=False).agg(Energia_kWh=("Energia_kWh","sum"))
            resultados.append({"Anio":int(periodo[:4]),"Mes":int(periodo[4:6]),"Proyecto":proyecto,"Generacion_MWh":df["Energia_kWh"].abs().sum()/1000,"Fuente":"API"})
    return pd.DataFrame(resultados,columns=["Anio","Mes","Proyecto","Generacion_MWh","Fuente"]), sorted(set(advertencias))


@st.cache_data(show_spinner=False)
def procesar_bess_api(periodos):
    mensuales, diarios, advertencias = [], [], []
    for periodo in periodos:
        registros, correctos = [], 0
        for mpid in BESS_MPIDS:
            datos, error = obtener_datos_prmte(periodo+"012345",mpid,3); regs = extraer_mediciones(datos,3)
            if datos is None or not regs: advertencias.append(f"BESS, {periodo_a_texto(periodo)}, {mpid}: {error or 'sin mediciones'}"); continue
            correctos += 1; registros.extend(regs)
        if correctos != len(BESS_MPIDS) or not registros: continue
        df = pd.DataFrame(registros); df["Fecha"] = parse_fecha_safe(df["Fecha"]); df["Energia_kWh"] = convertir_numerico(df["Energia_kWh"]).fillna(0)
        df = df.dropna(subset=["Fecha"]).groupby("Fecha",as_index=False).agg(Energia_kWh=("Energia_kWh","sum")); df["Energia_MWh"] = df["Energia_kWh"].abs()/1000
        df["Tipo"] = "PPA"; df.loc[(df["Fecha"].dt.hour>=6)&(df["Fecha"].dt.hour<21),"Tipo"] = "SPOT"; df["Fecha_Dia"] = df["Fecha"].dt.normalize()
        diario = df.groupby(["Fecha_Dia","Tipo"],as_index=False).agg(Energia_MWh=("Energia_MWh","sum")).pivot_table(index="Fecha_Dia",columns="Tipo",values="Energia_MWh",fill_value=0).reset_index(); diario.columns.name = None
        for c in ["PPA","SPOT"]:
            if c not in diario.columns: diario[c] = 0
        diario = diario.rename(columns={"PPA":"Energia_PPA_MWh","SPOT":"Energia_Spot_MWh"}); diario["Energia_Total_MWh"] = diario["Energia_PPA_MWh"] + diario["Energia_Spot_MWh"]; diario["Anio"] = diario["Fecha_Dia"].dt.year; diario["Mes"] = diario["Fecha_Dia"].dt.month; diarios.append(diario)
        cmg = descargar_cmg_mes(periodo)
        if cmg is None or cmg.empty: advertencias.append(f"BESS, {periodo_a_texto(periodo)}: CMG no disponible; energía disponible, ingresos Spot no valorizados"); df["CMG_USD_MWh"] = 0.0
        else: df = df.merge(cmg,on="Fecha",how="left"); df["CMG_USD_MWh"] = df["CMG_USD_MWh"].fillna(0)
        df["Precio"] = PRECIO_PPA; df.loc[df["Tipo"]=="SPOT","Precio"] = df["CMG_USD_MWh"]; df["Ingreso_USD"] = df["Energia_MWh"] * df["Precio"]
        eppa = df.loc[df["Tipo"]=="PPA","Energia_MWh"].sum(); espot = df.loc[df["Tipo"]=="SPOT","Energia_MWh"].sum(); ippa = df.loc[df["Tipo"]=="PPA","Ingreso_USD"].sum(); ispot = df.loc[df["Tipo"]=="SPOT","Ingreso_USD"].sum()
        mensuales.append({"Anio":int(periodo[:4]),"Mes":int(periodo[4:6]),"Energia_PPA_MWh":eppa,"Energia_Spot_MWh":espot,"Energia_Total_MWh":eppa+espot,"Ingreso_PPA_USD":ippa,"Ingreso_Spot_USD":ispot,"Ingreso_Total_USD":ippa+ispot,"Fuente":"API"})
    cols_m = ["Anio","Mes","Energia_PPA_MWh","Energia_Spot_MWh","Energia_Total_MWh","Ingreso_PPA_USD","Ingreso_Spot_USD","Ingreso_Total_USD","Fuente"]
    cols_d = ["Fecha_Dia","Energia_PPA_MWh","Energia_Spot_MWh","Energia_Total_MWh","Anio","Mes"]
    return pd.DataFrame(mensuales,columns=cols_m), (pd.concat(diarios,ignore_index=True) if diarios else pd.DataFrame(columns=cols_d)), sorted(set(advertencias))


def procesar_registros_diarios(registros, proyecto):
    df = pd.DataFrame(registros); df["Fecha"] = parse_fecha_safe(df["Fecha"]); df["Energia_kWh"] = convertir_numerico(df["Energia_kWh"]).fillna(0); df = df.dropna(subset=["Fecha"])
    if df.empty: return pd.DataFrame()
    df["Fecha_Dia"] = df["Fecha"].dt.normalize(); diario = df.groupby("Fecha_Dia",as_index=False).agg(Energia_MWh=("Energia_kWh",lambda s:s.abs().sum()/1000)); diario["Proyecto"] = proyecto; diario["Anio"] = diario["Fecha_Dia"].dt.year; diario["Mes"] = diario["Fecha_Dia"].dt.month; diario["Fuente"] = "API"
    return diario[["Fecha_Dia","Anio","Mes","Proyecto","Energia_MWh","Fuente"]]


@st.cache_data(show_spinner=False)
def descargar_diarios_carrusel(periodos, fecha_inicio, fecha_fin):
    resultados, advertencias = [], []; configuraciones = {"MARIA ELENA BESS":[{"mpid":m,"canal":3} for m in BESS_MPIDS]}; configuraciones.update(PROYECTOS)
    for proyecto in ORDEN_CARRUSEL:
        diarios = []
        for periodo in periodos:
            registros, correctos = [], 0
            for conf in configuraciones[proyecto]:
                datos, error = obtener_datos_prmte(periodo+"012345",conf["mpid"],conf["canal"]); regs = extraer_mediciones(datos,conf["canal"])
                if datos is None or not regs: advertencias.append(f"Carrusel {proyecto}, {periodo_a_texto(periodo)}, {conf['mpid']}: {error or 'sin mediciones'}"); continue
                correctos += 1; registros.extend(regs)
            if correctos == len(configuraciones[proyecto]) and registros:
                d = procesar_registros_diarios(registros,proyecto)
                if not d.empty: diarios.append(d)
        if not diarios: advertencias.append(f"Carrusel {proyecto}: sin datos diarios entre {fecha_inicio:%d-%m-%Y} y {fecha_fin:%d-%m-%Y}"); continue
        d = pd.concat(diarios,ignore_index=True); d = d[(d["Fecha_Dia"]>=fecha_inicio)&(d["Fecha_Dia"]<=fecha_fin)].copy()
        if d.empty: advertencias.append(f"Carrusel {proyecto}: sin mediciones en los últimos {DIAS_CARRUSEL} días"); continue
        resultados.append(d.groupby(["Fecha_Dia","Anio","Mes","Proyecto","Fuente"],as_index=False).agg(Energia_MWh=("Energia_MWh","sum")))
    cols = ["Fecha_Dia","Anio","Mes","Proyecto","Energia_MWh","Fuente"]
    return (pd.concat(resultados,ignore_index=True).sort_values(["Proyecto","Fecha_Dia"]) if resultados else pd.DataFrame(columns=cols)), sorted(set(advertencias))

# =========================================================
# DESCARGABLE PRMTE HORARIO
# =========================================================
def periodos_entre_fechas_diarias(fecha_inicio, fecha_fin):
    """Devuelve los meses YYYYMM necesarios para cubrir un rango diario."""
    inicio = pd.Timestamp(fecha_inicio).normalize()
    fin = pd.Timestamp(fecha_fin).normalize()
    return [p.strftime("%Y%m") for p in pd.period_range(inicio, fin, freq="M")]


def obtener_serie_prmte_rango(configuraciones, canal, nombre_columna, fecha_inicio, fecha_fin):
    """Descarga y suma los MPID de una instalación para un canal y rango."""
    registros = []
    periodos = periodos_entre_fechas_diarias(fecha_inicio, fecha_fin)

    for conf in configuraciones:
        mpid = conf["mpid"]
        canal_consulta = conf.get("canal", canal)
        for periodo in periodos:
            datos, error_api = obtener_datos_prmte(
                periodo + "012345",
                mpid,
                canal_consulta,
            )
            mediciones = extraer_mediciones(datos, canal_consulta)
            if datos is None or not mediciones:
                continue
            registros.extend(mediciones)

    if not registros:
        return pd.DataFrame(columns=["Fecha", nombre_columna])

    df = pd.DataFrame(registros)
    df["Fecha"] = parse_fecha_safe(df["Fecha"])
    df[nombre_columna] = convertir_numerico(df["Energia_kWh"]).fillna(0)
    df = df.dropna(subset=["Fecha"])

    inicio = pd.Timestamp(fecha_inicio).normalize()
    fin_exclusivo = pd.Timestamp(fecha_fin).normalize() + pd.Timedelta(days=1)
    df = df[(df["Fecha"] >= inicio) & (df["Fecha"] < fin_exclusivo)].copy()

    if df.empty:
        return pd.DataFrame(columns=["Fecha", nombre_columna])

    return (
        df.groupby("Fecha", as_index=False)[nombre_columna]
        .sum()
        .sort_values("Fecha")
    )


def completar_y_resamplear_horario(df, columna, fecha_inicio, fecha_fin):
    """Completa el calendario de 15 minutos y agrega la medición a una hora."""
    inicio = pd.Timestamp(fecha_inicio).normalize()
    fin_hora = pd.Timestamp(fecha_fin).normalize() + pd.Timedelta(hours=23)
    indice_horario = pd.date_range(inicio, fin_hora, freq="h")

    if df.empty:
        return pd.DataFrame({"Fecha": indice_horario, columna: pd.NA})

    horario = (
        df.set_index("Fecha")
        .resample("h")[columna]
        .sum(min_count=1)
        .reindex(indice_horario)
        .rename_axis("Fecha")
        .reset_index()
    )
    return horario


@st.cache_data(show_spinner=False)
def generar_excel_prmte_horario(fecha_inicio, fecha_fin):
    """
    Genera un Excel PRMTE horario para todos los parques.

    Incluye:
    - Inyección horaria, canal 3, para todos los parques.
    - María Elena PFV: Inyección y Neto PV Red según el neteo del ejemplo.
    - BESS María Elena: Inyección canal 3 y Carga canal 1.
    - Hoja Resumen con totales del rango.

    La caché no tiene TTL. El mismo rango no vuelve a consultar la API mientras
    la aplicación permanezca en ejecución.
    """
    fecha_inicio = pd.Timestamp(fecha_inicio).normalize()
    fecha_fin = pd.Timestamp(fecha_fin).normalize()
    if fecha_fin < fecha_inicio:
        raise ValueError("La fecha final no puede ser anterior a la fecha inicial.")

    configuraciones = {
        nombre: [dict(conf) for conf in lista]
        for nombre, lista in PROYECTOS.items()
    }
    configuraciones["BESS MARIA ELENA"] = [
        {"mpid": mpid, "canal": 3}
        for mpid in BESS_MPIDS
    ]

    # Series BESS necesarias tanto para su hoja como para el neteo de María Elena PFV.
    bess_inyeccion_15m = obtener_serie_prmte_rango(
        configuraciones["BESS MARIA ELENA"],
        3,
        "Inyección kWh",
        fecha_inicio,
        fecha_fin,
    )
    bess_carga_15m = obtener_serie_prmte_rango(
        [{"mpid": mpid, "canal": 1} for mpid in BESS_MPIDS],
        1,
        "Carga BESS kWh",
        fecha_inicio,
        fecha_fin,
    )

    hojas = {}
    resumen = []

    # Todos los parques solares.
    for proyecto in PROYECTOS:
        df_15m = obtener_serie_prmte_rango(
            configuraciones[proyecto],
            3,
            "Inyección kWh",
            fecha_inicio,
            fecha_fin,
        )

        if proyecto == "MARIA ELENA PFV":
            base = df_15m.merge(
                bess_inyeccion_15m.rename(columns={"Inyección kWh": "BESS Raw kWh"}),
                on="Fecha",
                how="left",
            )
            base["BESS Raw kWh"] = convertir_numerico(base["BESS Raw kWh"]).fillna(0)
            base["Neto PV Red kWh"] = base["Inyección kWh"]
            base.loc[base["BESS Raw kWh"] > 0, "Neto PV Red kWh"] = 0

            inyeccion_h = completar_y_resamplear_horario(
                base[["Fecha", "Inyección kWh"]],
                "Inyección kWh",
                fecha_inicio,
                fecha_fin,
            )
            neto_h = completar_y_resamplear_horario(
                base[["Fecha", "Neto PV Red kWh"]],
                "Neto PV Red kWh",
                fecha_inicio,
                fecha_fin,
            )
            horario = inyeccion_h.merge(neto_h, on="Fecha", how="outer")
            hojas[proyecto] = horario
            resumen.append({
                "Parque": proyecto,
                "Inyección MWh": horario["Inyección kWh"].sum(min_count=1) / 1000,
                "Neto PV Red MWh": horario["Neto PV Red kWh"].sum(min_count=1) / 1000,
                "Carga MWh": pd.NA,
            })
        else:
            horario = completar_y_resamplear_horario(
                df_15m,
                "Inyección kWh",
                fecha_inicio,
                fecha_fin,
            )
            hojas[proyecto] = horario
            resumen.append({
                "Parque": proyecto,
                "Inyección MWh": horario["Inyección kWh"].sum(min_count=1) / 1000,
                "Neto PV Red MWh": pd.NA,
                "Carga MWh": pd.NA,
            })

    # BESS: inyección y carga horaria.
    bess_h = completar_y_resamplear_horario(
        bess_inyeccion_15m,
        "Inyección kWh",
        fecha_inicio,
        fecha_fin,
    )
    carga_h = completar_y_resamplear_horario(
        bess_carga_15m,
        "Carga BESS kWh",
        fecha_inicio,
        fecha_fin,
    )
    bess_h = bess_h.merge(carga_h, on="Fecha", how="outer")
    hojas["BESS MARIA ELENA"] = bess_h
    resumen.append({
        "Parque": "BESS MARIA ELENA",
        "Inyección MWh": bess_h["Inyección kWh"].sum(min_count=1) / 1000,
        "Neto PV Red MWh": pd.NA,
        "Carga MWh": bess_h["Carga BESS kWh"].sum(min_count=1) / 1000,
    })

    df_resumen = pd.DataFrame(resumen)
    df_resumen.insert(0, "Fecha inicio", fecha_inicio.strftime("%Y-%m-%d"))
    df_resumen.insert(1, "Fecha fin", fecha_fin.strftime("%Y-%m-%d"))

    buffer = BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        df_resumen.to_excel(writer, sheet_name="Resumen", index=False)

        nombres_usados = {"Resumen"}
        for proyecto in ORDEN_CARRUSEL[1:] + ["BESS MARIA ELENA"]:
            if proyecto not in hojas:
                continue
            nombre_hoja = proyecto[:31]
            if nombre_hoja in nombres_usados:
                nombre_hoja = (nombre_hoja[:27] + "_PRMTE")[:31]
            nombres_usados.add(nombre_hoja)
            hojas[proyecto].to_excel(writer, sheet_name=nombre_hoja, index=False)

        # Formato general de cada hoja.
        for nombre_hoja, ws in writer.sheets.items():
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = ws.dimensions
            ws.sheet_view.showGridLines = False

            for celda in ws[1]:
                celda.fill = PatternFill("solid", fgColor="1565C0")
                celda.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
                celda.alignment = Alignment(horizontal="center", vertical="center")
            ws.row_dimensions[1].height = 22

            lado = Side(style="thin", color="DFE1E6")
            borde = Border(left=lado, right=lado, top=lado, bottom=lado)
            for fila in ws.iter_rows(min_row=2):
                for celda in fila:
                    celda.border = borde
                    celda.font = Font(name="Calibri", size=10, color="000000")
                    if celda.column == 1 and nombre_hoja != "Resumen":
                        celda.number_format = "yyyy-mm-dd hh:mm"
                    elif isinstance(celda.value, (int, float)):
                        celda.number_format = '#,##0.000'

            for columna in ws.columns:
                largo = max(len(str(c.value)) if c.value is not None else 0 for c in columna)
                ws.column_dimensions[columna[0].column_letter].width = min(max(largo + 2, 12), 28)

            ws.page_setup.orientation = "landscape"
            ws.page_setup.fitToWidth = 1
            ws.sheet_properties.pageSetUpPr.fitToPage = True

    buffer.seek(0)
    return buffer.getvalue()

# =========================================================
# NUEVO EXCEL MENSUAL SEGUN FORMATO SOLICITADO
# =========================================================
@st.cache_data(show_spinner=False)
def generar_excel_consolidado_mes(df_bess, df_proyectos, anio, mes):
    """Crea una hoja Consolidado Mes con generación real ya cargada en el dashboard."""
    periodo_texto = f"{anio:04d}-{mes:02d}"
    filas = []

    datos_proyectos = df_proyectos[(df_proyectos["Anio"]==anio)&(df_proyectos["Mes"]==mes)].set_index("Proyecto")
    for proyecto in ORDEN_DESCARGABLE_MENSUAL:
        valor = datos_proyectos.loc[proyecto,"Generacion_MWh"] if proyecto in datos_proyectos.index else None
        if isinstance(valor, pd.Series): valor = valor.sum()
        filas.append({"Mes":periodo_texto,"Parque":proyecto,"Total Inyección MWh":valor,"Distribución contrato":None})

    bess = df_bess[(df_bess["Anio"]==anio)&(df_bess["Mes"]==mes)]
    if not bess.empty:
        fila_bess = bess.iloc[-1]
        total_bess = fila_bess["Energia_Total_MWh"]
        energia_ppa = fila_bess["Energia_PPA_MWh"]
        energia_spot = fila_bess["Energia_Spot_MWh"]
    else:
        total_bess = energia_ppa = energia_spot = None

    filas.append({"Mes":periodo_texto,"Parque":"BESS MARIA ELENA (PPA)","Total Inyección MWh":total_bess,"Distribución contrato":energia_ppa})
    filas.append({"Mes":periodo_texto,"Parque":"BESS MARIA ELENA (SPOT)","Total Inyección MWh":None,"Distribución contrato":energia_spot})
    tabla = pd.DataFrame(filas)

    buffer = BytesIO()
    nombre_hoja = f"Consolidado {MESES_COMPLETOS[mes]}"[:31]
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        tabla.to_excel(writer, sheet_name=nombre_hoja, index=False)
        ws = writer.sheets[nombre_hoja]

        # Formato similar a la imagen.
        fill_header = PatternFill("solid", fgColor="FFFFFF")
        fill_valores = PatternFill("solid", fgColor="EBF1DE")
        lado = Side(style="dotted", color="000000")
        borde_punteado = Border(left=lado,right=lado,top=lado,bottom=lado)
        borde_header = Border(bottom=Side(style="thin",color="000000"))

        for c in ws[1]:
            c.font = Font(name="Calibri",size=11,bold=True,color="000000")
            c.fill = fill_header
            c.alignment = Alignment(horizontal="center",vertical="center")
            c.border = borde_header
        ws.row_dimensions[1].height = 20

        # Combinar total BESS en las dos filas, como en la imagen.
        fila_ppa = len(ORDEN_DESCARGABLE_MENSUAL) + 2
        fila_spot = fila_ppa + 1
        ws.merge_cells(start_row=fila_ppa,start_column=3,end_row=fila_spot,end_column=3)
        ws.cell(fila_ppa,3).value = total_bess

        for fila in range(2, ws.max_row + 1):
            for col in range(1,5):
                c = ws.cell(fila,col)
                c.font = Font(name="Calibri",size=11,color="000000")
                c.border = borde_punteado
                c.alignment = Alignment(horizontal="left" if col in [1,2] else "right",vertical="center")
                if col in [3,4]:
                    c.fill = fill_valores
                    c.number_format = '#,##0.0'
            ws.row_dimensions[fila].height = 19

        ws.column_dimensions["A"].width = 13
        ws.column_dimensions["B"].width = 25
        ws.column_dimensions["C"].width = 22
        ws.column_dimensions["D"].width = 23
        ws.freeze_panes = "A2"
        ws.sheet_view.showGridLines = True
        ws.auto_filter.ref = f"A1:D{ws.max_row}"
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth = 1
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.print_area = f"A1:D{ws.max_row}"

    buffer.seek(0)
    return buffer.getvalue()

# =========================================================
# EXCEL CONSOLIDADO ANUAL
# =========================================================
@st.cache_data(show_spinner=False)
def generar_excel_consolidado_anual(df_bess, df_proyectos, anio):
    """
    Genera un resumen anual por proyecto:
    - una columna por mes;
    - una columna con la suma anual;
    - utiliza exclusivamente los valores ya cargados en el dashboard.
    """
    orden_anual = [
        "MARIA ELENA BESS",
        "MARIA ELENA PFV",
        "SAN PEDRO III",
        "DOÑA CARMEN",
        "LA QUINTA",
        "LA PERLA",
        "LA HUERTA",
        "CHACAICO",
        "SANCLEMENTE",
        "TRILALEO",
        "COLLANCO",
    ]

    # Proyectos solares.
    datos_proyectos = df_proyectos[df_proyectos["Anio"] == anio][
        ["Proyecto", "Mes", "Generacion_MWh"]
    ].copy()

    # BESS total mensual, correspondiente a PPA + Spot.
    datos_bess = df_bess[df_bess["Anio"] == anio][
        ["Mes", "Energia_Total_MWh"]
    ].copy()
    datos_bess = datos_bess.rename(
        columns={"Energia_Total_MWh": "Generacion_MWh"}
    )
    datos_bess["Proyecto"] = "MARIA ELENA BESS"

    datos_anuales = pd.concat(
        [
            datos_bess[["Proyecto", "Mes", "Generacion_MWh"]],
            datos_proyectos,
        ],
        ignore_index=True,
    )

    tabla = (
        datos_anuales.pivot_table(
            index="Proyecto",
            columns="Mes",
            values="Generacion_MWh",
            aggfunc="sum",
        )
        .reindex(orden_anual)
        .reindex(columns=range(1, 13))
    )

    tabla.columns = [MESES_CORTOS[mes] for mes in range(1, 13)]
    tabla["Inyección Anual MWh"] = tabla.sum(axis=1, min_count=1)
    tabla = tabla.reset_index().rename(columns={"Proyecto": "Parque"})

    buffer = BytesIO()
    nombre_hoja = "Consolidado Anual"

    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        # Fila 1: título. Fila 2: encabezados.
        tabla.to_excel(
            writer,
            sheet_name=nombre_hoja,
            index=False,
            startrow=1,
        )

        ws = writer.sheets[nombre_hoja]

        # Título superior.
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=14)
        titulo = ws.cell(row=1, column=1)
        titulo.value = f"CONSOLIDADO ANUAL {anio} [MWh]"
        titulo.fill = PatternFill("solid", fgColor="0B3D91")
        titulo.font = Font(name="Calibri", size=13, bold=True, color="FFFFFF")
        titulo.alignment = Alignment(horizontal="left", vertical="center")
        ws.row_dimensions[1].height = 25

        # Encabezados.
        for col in range(1, 15):
            celda = ws.cell(row=2, column=col)
            celda.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
            celda.fill = PatternFill("solid", fgColor="1565C0")
            celda.alignment = Alignment(horizontal="center", vertical="center")
            lado = Side(style="thin", color="C9D1D9")
            celda.border = Border(left=lado, right=lado, top=lado, bottom=lado)
        ws.row_dimensions[2].height = 22

        # Cuerpo.
        fill_alterna = PatternFill("solid", fgColor="F5F7FA")
        fill_total = PatternFill("solid", fgColor="EBF1DE")
        lado = Side(style="thin", color="C9D1D9")
        borde = Border(left=lado, right=lado, top=lado, bottom=lado)

        for fila in range(3, ws.max_row + 1):
            for col in range(1, 15):
                celda = ws.cell(row=fila, column=col)
                celda.border = borde
                celda.font = Font(name="Calibri", size=11, color="000000")
                celda.alignment = Alignment(
                    horizontal="left" if col == 1 else "right",
                    vertical="center",
                )
                if fila % 2 == 0:
                    celda.fill = fill_alterna
                if col == 14:
                    celda.fill = fill_total
                    celda.font = Font(name="Calibri", size=11, bold=True, color="000000")
                if col >= 2:
                    celda.number_format = '#,##0.0'
            ws.row_dimensions[fila].height = 20

        # Fila de total del portafolio.
        fila_total = ws.max_row + 1
        ws.cell(fila_total, 1).value = "TOTAL PORTAFOLIO"
        for col in range(2, 15):
            letra = get_column_letter(col)
            ws.cell(fila_total, col).value = f"=SUM({letra}3:{letra}{fila_total - 1})"
            ws.cell(fila_total, col).number_format = '#,##0.0'

        for col in range(1, 15):
            celda = ws.cell(fila_total, col)
            celda.fill = PatternFill("solid", fgColor="DCE6F1")
            celda.font = Font(name="Calibri", size=11, bold=True, color="172B4D")
            celda.border = borde
            celda.alignment = Alignment(
                horizontal="left" if col == 1 else "right",
                vertical="center",
            )

        # Dimensiones y vista.
        ws.column_dimensions["A"].width = 25
        for col in range(2, 14):
            ws.column_dimensions[get_column_letter(col)].width = 12
        ws.column_dimensions["N"].width = 22
        ws.freeze_panes = "B3"
        ws.sheet_view.showGridLines = False
        ws.auto_filter.ref = f"A2:N{fila_total}"
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.print_area = f"A1:N{fila_total}"

    buffer.seek(0)
    return buffer.getvalue()

# =========================================================
# CARGA Y UNIONES
# =========================================================
try:
    df_budgets_bess = cargar_budgets_bess(); df_consolidado_bess = cargar_consolidado_bess(); df_budgets_proyectos = cargar_budgets_proyectos(); df_consolidado_proyectos = cargar_consolidado_proyectos()
except Exception as error:
    st.error("No fue posible validar budgets.xlsx."); st.info(str(error)); st.stop()

with st.spinner("Consultando meses pendientes de proyectos..."):
    df_proyectos_api, advertencias_proyectos = descargar_proyectos_pendientes(obtener_pendientes_proyectos(df_consolidado_proyectos))
fuentes_p = [df_consolidado_proyectos] + ([df_proyectos_api] if not df_proyectos_api.empty else [])
df_proyectos_final = pd.concat(fuentes_p,ignore_index=True); df_proyectos_final["Prioridad"] = df_proyectos_final["Fuente"].map({"API":1,"Excel consolidado":2}).fillna(2)
df_proyectos_final = df_proyectos_final.sort_values("Prioridad").drop_duplicates(["Anio","Mes","Proyecto"],keep="last").drop(columns="Prioridad").sort_values(["Proyecto","Anio","Mes"])

periodos_bess_excel = {f"{int(f.Anio):04d}{int(f.Mes):02d}" for f in df_consolidado_bess.itertuples()}; periodos_bess_pendientes = [p for p in PERIODOS_DISPONIBLES if p not in periodos_bess_excel]
with st.spinner("Consultando meses pendientes del BESS..."):
    df_bess_api, df_bess_diario_api, advertencias_bess = procesar_bess_api(periodos_bess_pendientes)
fuentes_b = [df_consolidado_bess] + ([df_bess_api] if not df_bess_api.empty else [])
df_bess_final = pd.concat(fuentes_b,ignore_index=True); df_bess_final["Prioridad"] = df_bess_final["Fuente"].map({"API":1,"Excel consolidado":2}).fillna(2)
df_bess_final = df_bess_final.sort_values("Prioridad").drop_duplicates(["Anio","Mes"],keep="last").drop(columns="Prioridad").sort_values(["Anio","Mes"])

with st.spinner("Preparando energía diaria para el carrusel..."):
    df_diarios_carrusel, advertencias_carrusel = descargar_diarios_carrusel(PERIODOS_CARRUSEL,FECHA_INICIO_CARRUSEL,FECHA_FIN_CARRUSEL)
advertencias = sorted(set(advertencias_proyectos + advertencias_bess + advertencias_carrusel))

# =========================================================
# SIDEBAR Y DESCARGABLE MENSUAL
# =========================================================
st.sidebar.header("Filtros")
st.sidebar.caption("Configuración de visualización")
anio_seleccionado = st.sidebar.selectbox("Año operacional",options=[ANIO_OPERACIONAL],index=0)
st.sidebar.divider()

# Meses disponibles: unión de BESS y parques, ordenados del más reciente al más antiguo.
periodos_reporte = sorted({(int(f.Anio),int(f.Mes)) for f in df_proyectos_final.itertuples()} | {(int(f.Anio),int(f.Mes)) for f in df_bess_final.itertuples()}, reverse=True)
periodo_descarga = st.sidebar.selectbox(
    "Mes del consolidado a descargar",
    periodos_reporte,
    format_func=lambda x: f"{MESES_COMPLETOS[x[1]]} {x[0]}",
    key="periodo_descargable_mensual",
)
anio_descarga, mes_descarga = periodo_descarga
archivo_mensual = generar_excel_consolidado_mes(df_bess_final,df_proyectos_final,anio_descarga,mes_descarga)

st.sidebar.download_button(
    "Descargar consolidado mensual",
    data=archivo_mensual,
    # Mes y año juntos, solo números: por ejemplo 092026.xlsx.
    file_name=f"{mes_descarga:02d}{anio_descarga:04d}.xlsx",
    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    use_container_width=True,
    key="descargar_consolidado_mensual",
)
st.sidebar.caption(f"Hoja: Consolidado {MESES_COMPLETOS[mes_descarga]}")

st.sidebar.divider()

# Descargable anual adicional.
archivo_anual = generar_excel_consolidado_anual(
    df_bess_final,
    df_proyectos_final,
    anio_seleccionado,
)

st.sidebar.download_button(
    "Descargar Consolidado Anual",
    data=archivo_anual,
    file_name=f"Consolidado Anual {anio_seleccionado}.xlsx",
    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    use_container_width=True,
    key="descargar_consolidado_anual",
)
st.sidebar.caption(f"Hoja: Consolidado Anual · Año {anio_seleccionado}")

st.sidebar.divider()
st.sidebar.subheader("Descarga PRMTE horario")

fecha_minima_prmte = pd.Timestamp(f"{ANIO_OPERACIONAL}-01-01").date()
fecha_maxima_prmte = pd.Timestamp.now().normalize().date()
fecha_inicio_prmte = st.sidebar.date_input(
    "Fecha inicial PRMTE",
    value=pd.Timestamp.now().normalize().replace(day=1).date(),
    min_value=fecha_minima_prmte,
    max_value=fecha_maxima_prmte,
    key="fecha_inicio_prmte",
)
fecha_fin_prmte = st.sidebar.date_input(
    "Fecha final PRMTE",
    value=fecha_maxima_prmte,
    min_value=fecha_minima_prmte,
    max_value=fecha_maxima_prmte,
    key="fecha_fin_prmte",
)

if "archivo_prmte_horario" not in st.session_state:
    st.session_state.archivo_prmte_horario = None
if "nombre_archivo_prmte" not in st.session_state:
    st.session_state.nombre_archivo_prmte = None
if "rango_archivo_prmte" not in st.session_state:
    st.session_state.rango_archivo_prmte = None

if st.sidebar.button(
    "Preparar archivo PRMTE",
    use_container_width=True,
    key="preparar_archivo_prmte",
):
    if fecha_fin_prmte < fecha_inicio_prmte:
        st.sidebar.error("La fecha final debe ser igual o posterior a la fecha inicial.")
    else:
        with st.spinner("Descargando y consolidando PRMTE horario..."):
            st.session_state.archivo_prmte_horario = generar_excel_prmte_horario(
                fecha_inicio_prmte,
                fecha_fin_prmte,
            )
            st.session_state.nombre_archivo_prmte = (
                f"PRMTE_{fecha_inicio_prmte:%Y%m%d}_{fecha_fin_prmte:%Y%m%d}.xlsx"
            )
            st.session_state.rango_archivo_prmte = (
                fecha_inicio_prmte,
                fecha_fin_prmte,
            )

if st.session_state.archivo_prmte_horario is not None:
    st.sidebar.download_button(
        "Descargar PRMTE horario",
        data=st.session_state.archivo_prmte_horario,
        file_name=st.session_state.nombre_archivo_prmte,
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        use_container_width=True,
        key="descargar_prmte_horario",
    )
    rango_prmte = st.session_state.rango_archivo_prmte
    st.sidebar.caption(
        f"Archivo preparado: {rango_prmte[0]:%d-%m-%Y} a {rango_prmte[1]:%d-%m-%Y}."
    )

st.sidebar.caption("Los consolidados usan los valores ya cargados; no vuelven a consultar la API.")
st.sidebar.caption("El archivo PRMTE consulta la API solo al presionar ‘Preparar archivo PRMTE’ para un rango nuevo.")
st.sidebar.caption("Los demás datos API se actualizan cuando la aplicación se reinicia.")

# =========================================================
# ENCABEZADO Y ADVERTENCIAS
# =========================================================
col_titulo,col_estado = st.columns([3.5,1],vertical_alignment="center")
with col_titulo:
    st.title("Dashboard de operación"); st.caption(f"BESS María Elena y portafolio de generación · Período {FECHA_INICIO} a {FECHA_FIN}")
with col_estado:
    st.markdown('<div style="background:#EAF2FD;color:#0B3D91;border:1px solid #B7D4F4;border-radius:10px;padding:10px 14px;text-align:center;font-size:.85rem;font-weight:650;">● Dashboard operativo</div>',unsafe_allow_html=True)
if advertencias:
    with st.expander(f"Estado de datos: {len(advertencias)} advertencias",expanded=False):
        st.warning("Algunas consultas no estuvieron disponibles. Los meses consolidados del Excel no se ven afectados.")
        for a in advertencias: st.markdown(f"- {a}")
else: st.success("Todos los datos requeridos fueron cargados correctamente.",icon="✅")

# =========================================================
# CARRUSEL
# =========================================================
st.header("Monitoreo diario de proyectos")
st.caption(f"Rotación automática cada {SEGUNDOS_ROTACION_GRAFICOS} segundos. Cada proyecto muestra los últimos {DIAS_CARRUSEL} días, desde {FECHA_INICIO_CARRUSEL:%d-%m-%Y} hasta {FECHA_FIN_CARRUSEL:%d-%m-%Y}.")
for llave,valor in {"carrusel_indice":0,"carrusel_pausado":False,"carrusel_inicializado":False,"carrusel_accion_manual":False}.items():
    if llave not in st.session_state: st.session_state[llave] = valor
proyectos_carrusel_disponibles = [p for p in ORDEN_CARRUSEL if p in set(df_diarios_carrusel["Proyecto"].unique())]

def carrusel_anterior():
    if proyectos_carrusel_disponibles: st.session_state.carrusel_indice = (st.session_state.carrusel_indice-1)%len(proyectos_carrusel_disponibles); st.session_state.carrusel_accion_manual = True

def carrusel_siguiente():
    if proyectos_carrusel_disponibles: st.session_state.carrusel_indice = (st.session_state.carrusel_indice+1)%len(proyectos_carrusel_disponibles); st.session_state.carrusel_accion_manual = True

def carrusel_pausar_reanudar():
    st.session_state.carrusel_pausado = not st.session_state.carrusel_pausado; st.session_state.carrusel_accion_manual = True

def carrusel_seleccionar():
    seleccionado = st.session_state.selector_carrusel
    if seleccionado in proyectos_carrusel_disponibles: st.session_state.carrusel_indice = proyectos_carrusel_disponibles.index(seleccionado)
    st.session_state.carrusel_accion_manual = True

@st.fragment(run_every=SEGUNDOS_ROTACION_GRAFICOS)
def mostrar_carrusel_diario():
    if not proyectos_carrusel_disponibles: st.info("No hay datos diarios disponibles para mostrar en el carrusel."); return
    cantidad = len(proyectos_carrusel_disponibles); st.session_state.carrusel_indice %= cantidad
    if not st.session_state.carrusel_inicializado: st.session_state.carrusel_inicializado = True
    elif st.session_state.carrusel_accion_manual: st.session_state.carrusel_accion_manual = False
    elif not st.session_state.carrusel_pausado: st.session_state.carrusel_indice = (st.session_state.carrusel_indice+1)%cantidad
    proyecto = proyectos_carrusel_disponibles[st.session_state.carrusel_indice]; st.session_state.selector_carrusel = proyecto
    c1,c2,c3,c4 = st.columns([1,1,1,3])
    with c1: st.button("◀ Anterior",use_container_width=True,key="boton_carrusel_anterior",on_click=carrusel_anterior)
    with c2: st.button("▶ Reanudar" if st.session_state.carrusel_pausado else "⏸ Pausar",use_container_width=True,key="boton_carrusel_pausa",on_click=carrusel_pausar_reanudar)
    with c3: st.button("Siguiente ▶",use_container_width=True,key="boton_carrusel_siguiente",on_click=carrusel_siguiente)
    with c4: st.selectbox("Ir al proyecto",proyectos_carrusel_disponibles,key="selector_carrusel",on_change=carrusel_seleccionar)
    proyecto = proyectos_carrusel_disponibles[st.session_state.carrusel_indice]
    datos = df_diarios_carrusel[df_diarios_carrusel["Proyecto"]==proyecto][["Fecha_Dia","Energia_MWh"]].copy(); calendario = pd.DataFrame({"Fecha_Dia":pd.date_range(FECHA_INICIO_CARRUSEL,FECHA_FIN_CARRUSEL,freq="D")}); datos = calendario.merge(datos,on="Fecha_Dia",how="left")
    fig = go.Figure(); fig.add_bar(x=datos["Fecha_Dia"],y=datos["Energia_MWh"],name="Energía diaria",marker_color=COLOR_PRIMARIO,width=.58*24*60*60*1000,hovertemplate="<b>%{x|%d-%m-%Y}</b><br>Energía: %{y:,.2f} MWh<extra></extra>")
    fig.update_layout(title=f"Energía diaria · {proyecto} · {FECHA_INICIO_CARRUSEL:%d-%m-%Y} a {FECHA_FIN_CARRUSEL:%d-%m-%Y}",xaxis_title="Fecha",yaxis_title="Energía [MWh]",showlegend=False); fig.update_xaxes(dtick=24*60*60*1000,tickformat="%d-%m",range=[FECHA_INICIO_CARRUSEL-pd.Timedelta(hours=12),FECHA_FIN_CARRUSEL+pd.Timedelta(hours=12)]); fig.update_yaxes(rangemode="tozero"); aplicar_estilo_grafico(fig,500,mostrar_leyenda=False,margen_inferior=60)
    clave = quitar_tildes(proyecto).lower().replace(" ","_"); st.plotly_chart(fig,use_container_width=True,key=f"carrusel_{clave}_{FECHA_INICIO_CARRUSEL:%Y%m%d}_{FECHA_FIN_CARRUSEL:%Y%m%d}")
    estado = "Pausado" if st.session_state.carrusel_pausado else f"Rotación automática cada {SEGUNDOS_ROTACION_GRAFICOS} segundos"; st.caption(f"Vista {st.session_state.carrusel_indice+1} de {cantidad} · {estado} · Fuente: API")
mostrar_carrusel_diario()

# =========================================================
# BESS
# =========================================================
st.header("BESS María Elena")
bess_anio = df_bess_final[df_bess_final["Anio"]==anio_seleccionado].copy()
if bess_anio.empty: st.info("No hay información disponible del BESS.")
else:
    meses_bess = sorted(bess_anio["Mes"].astype(int).unique().tolist(),reverse=True); mes_bess = st.selectbox("Mes para indicadores BESS",meses_bess,format_func=lambda m:MESES_COMPLETOS[m],key="mes_kpi_bess"); fila = bess_anio[bess_anio["Mes"]==mes_bess].iloc[0]; budget_mes = df_budgets_bess[(df_budgets_bess["Anio"]==anio_seleccionado)&(df_budgets_bess["Mes"]==mes_bess)]; diferencia = fila["Energia_Total_MWh"]-budget_mes.iloc[0]["Budget_Generacion_MWh"] if not budget_mes.empty else 0
    c1,c2,c3,c4 = st.columns(4); c1.metric("Inyección total del mes",f"{fila['Energia_Total_MWh']:,.1f} MWh",delta=f"{diferencia:+,.1f} MWh vs budget"); c2.metric("Energía PPA del mes",f"{fila['Energia_PPA_MWh']:,.1f} MWh"); c3.metric("Energía Spot del mes",f"{fila['Energia_Spot_MWh']:,.1f} MWh"); c4.metric("Ingreso del mes",f"USD {fila['Ingreso_Total_USD']:,.0f}"); st.caption(f"Fuente: {fila['Fuente']} · Precio PPA: USD {PRECIO_PPA:,.2f}/MWh")

st.subheader("Distribución diaria de energía")
if df_bess_diario_api.empty: st.info("No hay información diaria disponible desde la API. Los meses del Excel contienen información mensual.")
else:
    opciones = [(int(f.Anio),int(f.Mes)) for f in df_bess_diario_api[["Anio","Mes"]].drop_duplicates().sort_values(["Anio","Mes"],ascending=False).itertuples()]; ad,md = st.selectbox("Mes para gráfico diario",opciones,format_func=lambda x:f"{MESES_COMPLETOS[x[1]]} {x[0]}",key="periodo_diario_bess"); diario_mes = df_bess_diario_api[(df_bess_diario_api["Anio"]==ad)&(df_bess_diario_api["Mes"]==md)].copy(); primer = pd.Timestamp(ad,md,1); ultimo = pd.Timestamp(ad,md,calendar.monthrange(ad,md)[1]); hoy_n = pd.Timestamp.now().normalize()
    if ad == hoy_n.year and md == hoy_n.month: ultimo = min(ultimo,hoy_n)
    diario_mes = pd.DataFrame({"Fecha_Dia":pd.date_range(primer,ultimo,freq="D")}).merge(diario_mes[["Fecha_Dia","Energia_PPA_MWh","Energia_Spot_MWh","Energia_Total_MWh"]],on="Fecha_Dia",how="left").fillna(0); ancho = .40*24*60*60*1000; fig_d = go.Figure(); fig_d.add_bar(x=diario_mes["Fecha_Dia"],y=diario_mes["Energia_PPA_MWh"],name="PPA real",marker_color=COLOR_PRIMARIO,width=ancho); fig_d.add_bar(x=diario_mes["Fecha_Dia"],y=diario_mes["Energia_Spot_MWh"],name="Spot real",marker_color=COLOR_SECUNDARIO,width=ancho); fig_d.update_layout(title=f"Inyección diaria BESS · {MESES_COMPLETOS[md]} {ad}",barmode="stack",xaxis_title="Día del mes",yaxis_title="Energía [MWh]"); fig_d.update_xaxes(dtick=24*60*60*1000,tickformat="%d",range=[primer-pd.Timedelta(hours=12),ultimo+pd.Timedelta(hours=12)]); fig_d.update_yaxes(range=[0,120],tickmode="linear",tick0=0,dtick=20); aplicar_estilo_grafico(fig_d,560,leyenda_abajo=True,margen_inferior=115); st.plotly_chart(fig_d,use_container_width=True,key="grafico_diario_bess")

con_b = pd.DataFrame({"Mes":range(1,13)}).merge(bess_anio,on="Mes",how="left"); bud_b = df_budgets_bess[df_budgets_bess["Anio"]==anio_seleccionado]; con_b = con_b.merge(bud_b[["Mes","Budget_Generacion_MWh","Budget_PPA_MWh"]],on="Mes",how="left")
for c in ["Energia_PPA_MWh","Energia_Spot_MWh","Energia_Total_MWh","Ingreso_Total_USD"]:
    if c not in con_b.columns: con_b[c] = 0
    con_b[c] = con_b[c].fillna(0)
con_b["Nombre_Mes"] = con_b["Mes"].map(MESES_CORTOS); fig_b = go.Figure(); fig_b.add_bar(x=con_b["Nombre_Mes"],y=con_b["Energia_PPA_MWh"],name="PPA real",marker_color=COLOR_PRIMARIO,width=.58); fig_b.add_bar(x=con_b["Nombre_Mes"],y=con_b["Energia_Spot_MWh"],name="Spot real",marker_color=COLOR_SECUNDARIO,width=.58); fig_b.add_scatter(x=con_b["Nombre_Mes"],y=con_b["Budget_Generacion_MWh"],name="Budget generación",mode="lines+markers",line=dict(color=COLOR_BUDGET,width=3)); fig_b.add_scatter(x=con_b["Nombre_Mes"],y=con_b["Budget_PPA_MWh"],name="Budget PPA",mode="lines+markers",line=dict(color=COLOR_BUDGET_PPA,width=3)); fig_b.update_layout(title="Consolidado anual BESS",barmode="stack",xaxis_title="Mes",yaxis_title="Energía [MWh]"); fig_b.update_yaxes(tickmode="linear",tick0=0,dtick=500,rangemode="tozero"); aplicar_estilo_grafico(fig_b,640,leyenda_abajo=True,margen_inferior=120); st.plotly_chart(fig_b,use_container_width=True,key="grafico_consolidado_anual_bess")
with st.expander("Ver acumulados del BESS"):
    a1,a2,a3,a4 = st.columns(4); a1.metric("Inyección acumulada",f"{bess_anio['Energia_Total_MWh'].sum():,.1f} MWh"); a2.metric("PPA acumulado",f"{bess_anio['Energia_PPA_MWh'].sum():,.1f} MWh"); a3.metric("Spot acumulado",f"{bess_anio['Energia_Spot_MWh'].sum():,.1f} MWh"); a4.metric("Ingreso acumulado",f"USD {bess_anio['Ingreso_Total_USD'].sum():,.0f}")

# =========================================================
# PORTAFOLIO
# =========================================================
st.header("Portafolio de generación")
proyectos_anio = df_proyectos_final[df_proyectos_final["Anio"]==anio_seleccionado].copy(); budgets_proyectos_anio = df_budgets_proyectos[df_budgets_proyectos["Anio"]==anio_seleccionado].copy(); gen_total = proyectos_anio.groupby("Mes",as_index=False).agg(Generacion_Total_MWh=("Generacion_MWh","sum")); bud_total = budgets_proyectos_anio.groupby("Mes",as_index=False).agg(Budget_Total_MWh=("Budget_Generacion_MWh","sum")); total = pd.DataFrame({"Mes":range(1,13)}).merge(gen_total,on="Mes",how="left").merge(bud_total,on="Mes",how="left"); total["Generacion_Total_MWh"] = total["Generacion_Total_MWh"].fillna(0); total["Nombre_Mes"] = total["Mes"].map(MESES_CORTOS)
meses_con_gen = total[total["Generacion_Total_MWh"]>0]
if not meses_con_gen.empty:
    ultimo_mes = int(meses_con_gen["Mes"].max()); fila_total = total[total["Mes"]==ultimo_mes].iloc[0]; gen_acum = total.loc[total["Mes"]<=ultimo_mes,"Generacion_Total_MWh"].sum(); bud_acum = total.loc[total["Mes"]<=ultimo_mes,"Budget_Total_MWh"].fillna(0).sum(); t1,t2,t3,t4 = st.columns(4); t1.metric(f"Generación total {MESES_COMPLETOS[ultimo_mes]}",f"{fila_total['Generacion_Total_MWh']:,.1f} MWh"); t2.metric(f"Budget total {MESES_COMPLETOS[ultimo_mes]}",f"{fila_total['Budget_Total_MWh']:,.1f} MWh" if pd.notna(fila_total["Budget_Total_MWh"]) else "Sin budget"); t3.metric("Generación acumulada",f"{gen_acum:,.1f} MWh"); t4.metric("Diferencia acumulada",f"{gen_acum-bud_acum:+,.1f} MWh",delta=f"{gen_acum-bud_acum:+,.1f} MWh vs budget")
fig_t = go.Figure(); fig_t.add_bar(x=total["Nombre_Mes"],y=total["Generacion_Total_MWh"],name="Generación real total",marker_color=COLOR_PRIMARIO,width=.38,offsetgroup="real"); fig_t.add_bar(x=total["Nombre_Mes"],y=total["Budget_Total_MWh"],name="Budget total",marker_color=COLOR_BUDGET_PROYECTO,width=.38,offsetgroup="budget"); fig_t.update_layout(title="Generación mensual total vs budget",barmode="group",xaxis_title="Mes",yaxis_title="Generación [MWh]"); aplicar_estilo_grafico(fig_t,560,leyenda_abajo=True,margen_inferior=110); st.plotly_chart(fig_t,use_container_width=True,key="grafico_generacion_total_proyectos")

st.subheader("Detalle mensual por proyecto"); mes_actual_operacional = int(FECHA_FIN[5:7]); columnas = st.columns(2)
for i,elemento in enumerate(ORDEN_GRAFICOS):
    if elemento == "CONSOLIDADO PMGDS": datos = proyectos_anio[proyectos_anio["Proyecto"].isin(PROYECTOS_PMGD)].groupby("Mes",as_index=False).agg(Generacion_MWh=("Generacion_MWh","sum")); bud = budgets_proyectos_anio[budgets_proyectos_anio["Proyecto"].isin(PROYECTOS_PMGD)].groupby("Mes",as_index=False).agg(Budget_Generacion_MWh=("Budget_Generacion_MWh","sum")); titulo = "Consolidado PMGDs"
    else: datos = proyectos_anio[proyectos_anio["Proyecto"]==elemento][["Mes","Generacion_MWh"]]; bud = budgets_proyectos_anio[budgets_proyectos_anio["Proyecto"]==elemento][["Mes","Budget_Generacion_MWh"]]; titulo = elemento
    con = pd.DataFrame({"Mes":range(1,13)}).merge(datos,on="Mes",how="left").merge(bud,on="Mes",how="left"); con = con[con["Mes"]<=mes_actual_operacional].copy(); con["Generacion_MWh"] = con["Generacion_MWh"].fillna(0); con["Nombre_Mes"] = con["Mes"].map(MESES_CORTOS); fig = go.Figure(); fig.add_bar(x=con["Nombre_Mes"],y=con["Generacion_MWh"],name="Generación",marker_color=COLOR_PRIMARIO,width=.52); fig.add_scatter(x=con["Nombre_Mes"],y=con["Budget_Generacion_MWh"],name="Budget",mode="lines+markers",line=dict(color=COLOR_BUDGET_PROYECTO,width=2),marker=dict(size=5)); fig.update_layout(title=titulo,yaxis_title="MWh",bargap=.38); aplicar_estilo_grafico(fig,340,mostrar_leyenda=(i==0),margen_inferior=45); clave = quitar_tildes(elemento).lower().replace(" ","_")
    with columnas[i%2]: st.plotly_chart(fig,use_container_width=True,key=f"grafico_{clave}")

with st.expander("Fuentes utilizadas por proyecto",expanded=False):
    tabla = proyectos_anio[["Proyecto","Mes","Fuente"]].sort_values(["Proyecto","Mes"]).copy(); tabla["Mes"] = tabla["Mes"].map(MESES_COMPLETOS); st.dataframe(tabla,use_container_width=True,hide_index=True)
st.divider(); st.caption(f"Inicio de la aplicación: {HORA_INICIO_APLICACION:%d-%m-%Y %H:%M:%S} · Las API se actualizan al reiniciar la aplicación")
